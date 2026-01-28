"""
ONI V24 Module Component
Part of the ONI Automation Framework.
Managed by The Jewel Polishing Protocol.
"""

"""
============================================================================
ONI_Excel.py - Python Integration for Excel Badge Generation
Version: 1.0
Description: Generate badges in Excel using openpyxl (no Excel required!)
Requirements: pip install openpyxl Pillow
============================================================================
"""

import random
import hashlib
from datetime import datetime
from typing import List, Dict, Tuple, Optional
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ============================================================================
# SEED MANAGEMENT
# ============================================================================

class ONISeedManager:
    """Manages random seeds"""
    
    _current_seed: Optional[str] = None
    _seed_history: List[str] = []
    
    @classmethod
    def generate_seed(cls, style_prefix: str = "XLS") -> str:
        """Generate new ONI seed"""
        timestamp = datetime.now().strftime("%Y%m%d")
        random_data = f"{datetime.now().isoformat()}{random.random()}"
        hash_obj = hashlib.sha256(random_data.encode())
        hash_hex = hash_obj.hexdigest()[:4].upper()
        
        seed = f"ONI-{style_prefix}-{timestamp}-{hash_hex}"
        cls._current_seed = seed
        cls._seed_history.append(seed)
        
        return seed
    
    @classmethod
    def set_seed(cls, seed: str) -> None:
        """Set current seed"""
        cls._current_seed = seed
        parts = seed.split("-")
        if len(parts) >= 4:
            seed_int = int(parts[3], 16)
            random.seed(seed_int)


# ============================================================================
# COLOR MANAGEMENT
# ============================================================================

class ONIColorManager:
    """Manage color palettes"""
    
    PALETTES = {
        "cyberpunk_v1": {
            "primary": ["00F0FF", "FF0055", "B026FF", "39FF14"],
            "accent": ["FFFF00", "00FF88", "FF6B00"],
            "background": ["0A0A0F", "1A1A20", "101015"]
        },
        "minimalism_v1": {
            "primary": ["0052A5", "4D4D4D", "FFFFFF"],
            "accent": ["0072CE", "D9D9D9"],
            "background": ["F5F5F5", "FFFFFF"]
        },
        "vaporwave_v1": {
            "primary": ["FF71CE", "01CDFE", "FF0080"],
            "accent": ["B967FF", "FFC87C"],
            "background": ["2E173D", "0D0221"]
        }
    }
    
    @staticmethod
    def get_random_color(palette: str = "cyberpunk_v1", 
                        category: str = "primary") -> str:
        """Get random color from palette"""
        if palette not in ONIColorManager.PALETTES:
            return "FFFFFF"
        
        colors = ONIColorManager.PALETTES[palette].get(category, ["FFFFFF"])
        return random.choice(colors)


# ============================================================================
# BADGE GENERATOR
# ============================================================================

class ONIBadgeGenerator:
    """Generate badges in Excel"""
    
    def __init__(self):
        self.seed_manager = ONISeedManager()
        self.color_manager = ONIColorManager()
    
    def create_cyberpunk_badge(self, 
                              name: str = "AGENT NAME",
                              title: str = "SECURITY",
                              badge_id: str = "A7F3-9182",
                              start_row: int = 2,
                              start_col: int = 2) -> Tuple[Workbook, str]:
        """
        Create cyberpunk-styled badge
        
        Returns:
            Tuple of (Workbook, seed)
        """
        
        # Generate seed
        seed = self.seed_manager.generate_seed("CYB")
        
        # Create workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "ONI_Badge_Cyberpunk"
        
        # Get colors
        primary = self.color_manager.get_random_color("cyberpunk_v1", "primary")
        accent = self.color_manager.get_random_color("cyberpunk_v1", "accent")
        bg = "0A0A0F"
        
        # Set column widths
        for col in range(start_col, start_col + 6):
            ws.column_dimensions[get_column_letter(col)].width = 12
        
        # Set row heights
        for row in range(start_row, start_row + 10):
            ws.row_dimensions[row].height = 25
        
        # Main badge border
        self._apply_border(ws, start_row, start_col, 10, 6, primary, thick=True)
        self._apply_background(ws, start_row, start_col, 10, 6, bg)
        
        # Header
        header_cell = ws.cell(start_row, start_col)
        self._merge_cells(ws, start_row, start_col, 2, 6)
        header_cell.value = "◢ CYBERPUNK ID ◣"
        header_cell.font = Font(name="Consolas", size=14, bold=True, color="0A0A0F")
        header_cell.fill = PatternFill(start_color=primary, end_color=primary, fill_type="solid")
        header_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Photo area
        photo_cell = ws.cell(start_row + 2, start_col)
        self._merge_cells(ws, start_row + 2, start_col, 4, 2)
        photo_cell.value = "█▓▒░ PHOTO ░▒▓█"
        photo_cell.font = Font(name="Consolas", size=10, color=primary)
        photo_cell.fill = PatternFill(start_color="1A1A20", end_color="1A1A20", fill_type="solid")
        photo_cell.alignment = Alignment(horizontal="center", vertical="center")
        self._apply_border(ws, start_row + 2, start_col, 4, 2, accent)
        
        # Name label
        name_label = ws.cell(start_row + 2, start_col + 3)
        name_label.value = "▸ NAME"
        name_label.font = Font(size=9, bold=True, color=accent)
        
        # Name value
        name_cell = ws.cell(start_row + 3, start_col + 2)
        self._merge_cells(ws, start_row + 3, start_col + 2, 1, 4)
        name_cell.value = name.upper()
        name_cell.font = Font(size=12, bold=True, color=primary)
        name_cell.alignment = Alignment(horizontal="center")
        
        # Title label
        title_label = ws.cell(start_row + 4, start_col + 3)
        title_label.value = "▸ CLEARANCE"
        title_label.font = Font(size=8, bold=True, color=accent)
        
        # Title value
        title_cell = ws.cell(start_row + 5, start_col + 2)
        self._merge_cells(ws, start_row + 5, start_col + 2, 1, 4)
        title_cell.value = title.upper()
        title_cell.font = Font(size=10, color="FFFFFF")
        title_cell.alignment = Alignment(horizontal="center")
        
        # ID
        id_cell = ws.cell(start_row + 6, start_col)
        self._merge_cells(ws, start_row + 6, start_col, 1, 6)
        id_cell.value = f"━━━ ID: {badge_id} ━━━"
        id_cell.font = Font(name="Courier New", size=10, color=accent)
        id_cell.fill = PatternFill(start_color="101015", end_color="101015", fill_type="solid")
        id_cell.alignment = Alignment(horizontal="center")
        
        # Greebles (random decorative elements)
        greeble_chars = ["▪", "▫", "▸", "◂", "●", "○"]
        for _ in range(3):
            g_row = start_row + 7 + random.randint(0, 1)
            g_col = start_col + random.randint(0, 5)
            greeble = ws.cell(g_row, g_col)
            greeble.value = random.choice(greeble_chars)
            greeble.font = Font(size=8, color=accent)
        
        # Barcode
        barcode_cell = ws.cell(start_row + 8, start_col)
        self._merge_cells(ws, start_row + 8, start_col, 1, 6)
        barcode_cell.value = "║││║│││║║││║││║│║││"
        barcode_cell.font = Font(name="Courier New", size=14, color="000000")
        barcode_cell.fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        barcode_cell.alignment = Alignment(horizontal="center")
        
        # Footer
        footer_left = ws.cell(start_row + 9, start_col)
        footer_left.value = f"SEED: {seed}"
        footer_left.font = Font(size=7, color="808080")
        
        footer_right = ws.cell(start_row + 9, start_col + 4)
        footer_right.value = datetime.now().strftime("%m/%Y")
        footer_right.font = Font(size=7, color="808080")
        footer_right.alignment = Alignment(horizontal="right")
        
        print(f"✓ Cyberpunk badge generated!")
        print(f"  Name: {name}")
        print(f"  Title: {title}")
        print(f"  ID: {badge_id}")
        print(f"  Seed: {seed}")
        
        return wb, seed
    
    def create_minimal_badge(self,
                           name: str = "EMPLOYEE NAME",
                           title: str = "POSITION",
                           badge_id: str = "EMP-001",
                           start_row: int = 2,
                           start_col: int = 2) -> Tuple[Workbook, str]:
        """Create minimal corporate badge"""
        
        seed = self.seed_manager.generate_seed("MIN")
        
        wb = Workbook()
        ws = wb.active
        ws.title = "ONI_Badge_Minimal"
        
        corporate_blue = "0052A5"
        light_gray = "F5F5F5"
        
        # Set dimensions
        for col in range(start_col, start_col + 6):
            ws.column_dimensions[get_column_letter(col)].width = 12
        
        for row in range(start_row, start_row + 10):
            ws.row_dimensions[row].height = 25
        
        # Border
        self._apply_border(ws, start_row, start_col, 10, 6, "C8C8C8")
        self._apply_background(ws, start_row, start_col, 10, 6, "FFFFFF")
        
        # Header
        header_cell = ws.cell(start_row, start_col)
        self._merge_cells(ws, start_row, start_col, 2, 6)
        header_cell.value = "EMPLOYEE BADGE"
        header_cell.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
        header_cell.fill = PatternFill(start_color=corporate_blue, end_color=corporate_blue, fill_type="solid")
        header_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Photo
        photo_cell = ws.cell(start_row + 2, start_col)
        self._merge_cells(ws, start_row + 2, start_col, 4, 2)
        photo_cell.value = "PHOTO"
        photo_cell.font = Font(color="969696")
        photo_cell.fill = PatternFill(start_color=light_gray, end_color=light_gray, fill_type="solid")
        photo_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Name
        ws.cell(start_row + 2, start_col + 3).value = "Name:"
        ws.cell(start_row + 2, start_col + 3).font = Font(bold=True)
        ws.cell(start_row + 3, start_col + 3).value = name
        ws.cell(start_row + 3, start_col + 3).font = Font(size=11, bold=True)
        
        # Title
        ws.cell(start_row + 4, start_col + 3).value = "Title:"
        ws.cell(start_row + 4, start_col + 3).font = Font(bold=True)
        ws.cell(start_row + 5, start_col + 3).value = title
        
        # ID
        id_cell = ws.cell(start_row + 6, start_col)
        self._merge_cells(ws, start_row + 6, start_col, 1, 6)
        id_cell.value = f"ID: {badge_id}"
        id_cell.fill = PatternFill(start_color=light_gray, end_color=light_gray, fill_type="solid")
        id_cell.alignment = Alignment(horizontal="center")
        
        # Barcode
        barcode_cell = ws.cell(start_row + 7, start_col)
        self._merge_cells(ws, start_row + 7, start_col, 2, 6)
        barcode_cell.value = "| | || ||| | || | ||| ||"
        barcode_cell.font = Font(name="Courier New", size=16)
        barcode_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Footer
        footer = ws.cell(start_row + 9, start_col)
        footer.value = f"Valid: {datetime.now().strftime('%m/%d/%Y')}"
        footer.font = Font(size=8)
        
        print(f"✓ Minimal badge generated! Seed: {seed}")
        
        return wb, seed
    
    # Helper methods
    def _merge_cells(self, ws, start_row, start_col, rows, cols):
        """Merge cells"""
        end_row = start_row + rows - 1
        end_col = start_col + cols - 1
        ws.merge_cells(
            start_row=start_row, start_column=start_col,
            end_row=end_row, end_column=end_col
        )
    
    def _apply_border(self, ws, start_row, start_col, rows, cols, color, thick=False):
        """Apply border to range"""
        style = Side(style="thick" if thick else "thin", color=color)
        
        for r in range(start_row, start_row + rows):
            for c in range(start_col, start_col + cols):
                cell = ws.cell(r, c)
                cell.border = Border(
                    left=style if c == start_col else None,
                    right=style if c == start_col + cols - 1 else None,
                    top=style if r == start_row else None,
                    bottom=style if r == start_row + rows - 1 else None
                )
    
    def _apply_background(self, ws, start_row, start_col, rows, cols, color):
        """Apply background color to range"""
        fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        
        for r in range(start_row, start_row + rows):
            for c in range(start_col, start_col + cols):
                ws.cell(r, c).fill = fill


# ============================================================================
# BATCH PROCESSING
# ============================================================================

def batch_generate_badges(csv_path: str, 
                         output_folder: str = "output",
                         style: str = "cyberpunk"):
    """Generate badges from CSV file"""
    import csv
    import os
    
    if not os.path.exists(output_folder):
        os.makedirs(output_folder)
    
    generator = ONIBadgeGenerator()
    
    with open(csv_path, 'r') as f:
        reader = csv.DictReader(f)
        for i, row in enumerate(reader):
            name = row['Name']
            title = row['Title']
            badge_id = row['ID']
            
            if style == "cyberpunk":
                wb, seed = generator.create_cyberpunk_badge(name, title, badge_id)
            else:
                wb, seed = generator.create_minimal_badge(name, title, badge_id)
            
            output_path = os.path.join(output_folder, f"{badge_id}_{name.replace(' ', '_')}.xlsx")
            wb.save(output_path)
            
            print(f"  ✓ Saved: {output_path}")
    
    print(f"\n✓ Batch complete!")


# ============================================================================
# USAGE EXAMPLES
# ============================================================================

if __name__ == "__main__":
    print("=== ONI Excel Badge Generator ===\n")
    
    # Example 1: Single cyberpunk badge
    generator = ONIBadgeGenerator()
    wb, seed = generator.create_cyberpunk_badge(
        name="ALEX MERCER",
        title="CYBER SECURITY",
        badge_id="CS-4729"
    )
    wb.save("badge_cyberpunk.xlsx")
    print("Saved: badge_cyberpunk.xlsx\n")
    
    # Example 2: Minimal badge
    wb2, seed2 = generator.create_minimal_badge(
        name="JOHN SMITH",
        title="PROJECT MANAGER",
        badge_id="PM-001"
    )
    wb2.save("badge_minimal.xlsx")
    print("Saved: badge_minimal.xlsx\n")
    
    # Example 3: Batch generation (uncomment to use)
    # batch_generate_badges("employees.csv", "output", "cyberpunk")
    
    print("\n✓ All examples completed!")

